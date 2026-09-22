"""Generate the production-grade IBM HR Analytics capstone report.

Run with:
    .\.venv\Scripts\python.exe generate_capstone_report.py

The script writes Himanshu_Pathak_IBM_HR_Project_Report.docx in the project root.
"""

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path("Himanshu_Pathak_IBM_HR_Project_Report.docx")
NAVY = "17365D"
SLATE = "44546A"
PALE_BLUE = "EAF0F7"
PALE_TEAL = "E7F2F0"
PALE_GOLD = "FFF3D6"
WHITE = "FFFFFF"
BORDER = "9AA9B8"


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def set_cell_borders(cell, color: str = BORDER, size: str = "8") -> None:
    properties = cell._tc.get_or_add_tcPr()
    borders = properties.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_text(cell, text: str, bold: bool = False, color: str = NAVY, size: int = 10) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None, fill: str = PALE_BLUE) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        set_cell_shading(cell, NAVY)
        set_cell_borders(cell)
        set_cell_text(cell, header, bold=True, color=WHITE, size=9)
        if widths:
            cell.width = Inches(widths[index])
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            set_cell_shading(cells[index], fill if len(table.rows) % 2 == 0 else WHITE)
            set_cell_borders(cells[index])
            set_cell_text(cells[index], str(value), size=9)
            if widths:
                cells[index].width = Inches(widths[index])
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_callout(doc: Document, label: str, text: str, fill: str = PALE_TEAL) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.cell(0, 0)
    cell.width = Inches(6.45)
    set_cell_shading(cell, fill)
    set_cell_borders(cell, color=NAVY, size="12")
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run(label + "\n")
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor.from_string(NAVY)
    body = paragraph.add_run(text)
    body.font.name = "Times New Roman"
    body.font.size = Pt(10)
    body.font.color.rgb = RGBColor.from_string(SLATE)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_screenshot_placeholder(doc: Document, text: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.cell(0, 0)
    cell.width = Inches(6.45)
    set_cell_shading(cell, "EDEDED")
    set_cell_borders(cell, color=NAVY, size="14")
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(35)
    paragraph.paragraph_format.space_after = Pt(35)
    run = paragraph.add_run(text)
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor.from_string(NAVY)
    doc.add_paragraph("Replace this placeholder with the corresponding screenshot from the running Streamlit dashboard.")


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    paragraph = doc.add_heading(text, level=level)
    paragraph.paragraph_format.keep_with_next = True
    for run in paragraph.runs:
        run.font.name = "Times New Roman"
        run.font.color.rgb = RGBColor.from_string(NAVY if level == 1 else SLATE)
    return paragraph


def add_body(doc: Document, text: str, bold_prefix: str | None = None) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.line_spacing = 1.15
    paragraph.paragraph_format.space_after = Pt(7)
    if bold_prefix and text.startswith(bold_prefix):
        first = paragraph.add_run(bold_prefix)
        first.bold = True
        first.font.name = "Times New Roman"
        first.font.size = Pt(12)
        rest = paragraph.add_run(text[len(bold_prefix):])
        rest.font.name = "Times New Roman"
        rest.font.size = Pt(12)
    else:
        run = paragraph.add_run(text)
        run.font.name = "Times New Roman"
        run.font.size = Pt(12)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.line_spacing = 1.15
        paragraph.paragraph_format.space_after = Pt(4)
        run = paragraph.add_run(item)
        run.font.name = "Times New Roman"
        run.font.size = Pt(12)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("People Signals: IBM HR Analytics Platform | ")
    run.font.name = "Times New Roman"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(SLATE)
    field_begin = OxmlElement("w:fldChar")
    field_begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    field_end = OxmlElement("w:fldChar")
    field_end.set(qn("w:fldCharType"), "end")
    run._r.append(field_begin)
    run._r.append(instruction)
    run._r.append(field_end)


def page_break(doc: Document) -> None:
    doc.add_page_break()


def build_report() -> Path:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(7)

    for style_name, size, color in (("Title", 26, NAVY), ("Heading 1", 18, NAVY), ("Heading 2", 14, SLATE), ("Subtitle", 14, SLATE)):
        style = doc.styles[style_name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = style_name != "Subtitle"
    doc.styles["Heading 1"].paragraph_format.space_before = Pt(8)
    doc.styles["Heading 1"].paragraph_format.space_after = Pt(7)
    doc.styles["Heading 2"].paragraph_format.space_before = Pt(6)
    doc.styles["Heading 2"].paragraph_format.space_after = Pt(4)

    footer = section.footer.paragraphs[0]
    add_page_number(footer)

    # Page 1: formal cover.
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(80)
    run = paragraph.add_run("IBM SkillsBuild | BharatCares (CSRBOX)")
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(15)
    run.font.color.rgb = RGBColor.from_string(NAVY)
    doc.add_paragraph()
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(40)
    run = title.add_run("PEOPLE SIGNALS: IBM HR ANALYTICS PLATFORM")
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(25)
    run.font.color.rgb = RGBColor.from_string(NAVY)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_before = Pt(18)
    run = subtitle.add_run("A 5-Level Business Intelligence Decision System for Employee Attrition, Tenure Velocity, and Retention Planning")
    run.font.name = "Times New Roman"
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor.from_string(SLATE)
    doc.add_paragraph("\n\n")
    cover_table = doc.add_table(rows=4, cols=2)
    cover_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cover_table.autofit = False
    cover_rows = [
        ("Prepared By", "Himanshu Pathak"),
        ("Domain", "Data Analytics with AI"),
        ("Internship", "IBM SkillsBuild Data Analytics with AI Internship 2026"),
        ("Date", "September 2026"),
    ]
    for row, (label, value) in zip(cover_table.rows, cover_rows):
        set_cell_shading(row.cells[0], PALE_BLUE)
        set_cell_shading(row.cells[1], WHITE)
        set_cell_borders(row.cells[0], color=BORDER)
        set_cell_borders(row.cells[1], color=BORDER)
        set_cell_text(row.cells[0], label, bold=True, size=11)
        set_cell_text(row.cells[1], value, size=11)
        row.cells[0].width = Inches(1.7)
        row.cells[1].width = Inches(4.6)
    doc.add_paragraph("\n\n")
    add_callout(doc, "PROJECT SCOPE", "A production-oriented Streamlit platform that connects descriptive workforce analytics, supervised learning, employee segmentation, and evidence-linked HR actions.", PALE_TEAL)
    page_break(doc)

    # Page 2: originality and acknowledgements.
    add_heading(doc, "Certificate of Originality & Acknowledgments")
    add_body(doc, "I, Himanshu Pathak, declare that this capstone report and the accompanying People Signals: IBM HR Analytics Platform were prepared as part of the IBM SkillsBuild Data Analytics with AI Internship 2026. The analysis, documentation, visual design, and implementation decisions are presented for academic and professional learning purposes. External datasets and open-source libraries are acknowledged in the references section.")
    add_body(doc, "I understand that the employee-level risk outputs in this project are analytical prioritisation signals only. They must not be used as automated employment decisions and should be reviewed with appropriate human, ethical, privacy, and organisational safeguards.")
    doc.add_paragraph("\n")
    signature = doc.add_table(rows=2, cols=2)
    signature.alignment = WD_TABLE_ALIGNMENT.CENTER
    signature.autofit = False
    for row in signature.rows:
        for cell in row.cells:
            set_cell_borders(cell, color=WHITE, size="0")
    set_cell_text(signature.cell(0, 0), "Signature: __________________________", size=11)
    set_cell_text(signature.cell(0, 1), "Date: __________________", size=11)
    set_cell_text(signature.cell(1, 0), "Himanshu Pathak", size=11)
    set_cell_text(signature.cell(1, 1), "September 2026", size=11)
    add_heading(doc, "Acknowledgments", level=2)
    add_body(doc, "I sincerely thank IBM SkillsBuild for the structured learning environment and practical orientation toward data analytics with AI. I am grateful to BharatCares and CSRBOX for facilitating the internship experience and to mentors Kartik Hooda and Himanshu Souda for their guidance, feedback, and encouragement. I also acknowledge the maintainers and contributors of Python, pandas, Streamlit, Plotly, scikit-learn, KaggleHub, ReportLab, and python-docx, whose open-source work made this capstone possible.")
    add_callout(doc, "DOCUMENT NOTE", "This report is intentionally written as a complete capstone artifact. Bracketed screenshot boxes are reserved for final dashboard captures from the running application.", PALE_GOLD)
    page_break(doc)

    # Page 3: executive summary and framework.
    add_heading(doc, "Executive Summary & 5-Level BI Hierarchy Synthesis")
    add_body(doc, "People Signals: IBM HR Analytics Platform turns 1,470 employee records into a decision surface for retention, mobility, and workforce planning. The baseline workforce view shows a 16.1% attrition rate and average monthly income of $6,503. The system progresses from descriptive KPIs to lifecycle trends, diagnostic drivers, model-led risk and employee segments, and finally to strategic actions that are tied to measured evidence.")
    add_heading(doc, "Headline findings", level=2)
    add_bullets(doc, [
        "OverTime is a material retention lever: the observed attrition rate reaches 30.5% for the highest overtime group.",
        "Sales Representative turnover is 39.8%, making role-specific coaching and career-path interventions a priority.",
        "Early tenure is a clear lifecycle pressure point: the 0-1 year bucket records 34.9% attrition.",
        "The platform uses an optimal classification threshold of 0.220 so risk review reflects the cost of missing at-risk employees rather than relying on the default 0.50 cutoff.",
    ])
    add_heading(doc, "5-level synthesis", level=2)
    add_table(doc, ["Level", "Business question", "Decision output"], [
        ["1 / KPIs", "What is happening?", "Scale, attrition, income, departments, and roles."],
        ["2 / Trends", "Direction kya hai?", "Chronological tenure lifecycle and income progression."],
        ["3 / Drivers", "Why is it happening?", "Work design, role, distance, marital status, and experience diagnostics."],
        ["4 / Risks", "Who needs attention?", "Permutation importance, thresholded risk scores, and KMeans segments."],
        ["5 / Actions", "What should HR do?", "Prioritised interventions with measurable follow-through."],
    ], [1.05, 1.75, 3.65])
    page_break(doc)

    # Page 4: architecture and data quality.
    add_heading(doc, "Dataset Architecture, Schema & Data Quality Audit")
    add_body(doc, "The source is the Kaggle IBM HR Analytics Employee Attrition & Performance dataset. It contains 1,470 rows and 35 source attributes. The application downloads the dataset through KaggleHub inside the same single-file Streamlit app, then applies deterministic preparation before analysis and modelling.")
    add_heading(doc, "Analytical schema", level=2)
    add_table(doc, ["Attribute family", "Representative fields", "Analytical role"], [
        ["Work context", "Department, JobRole, BusinessTravel, OverTime", "Driver diagnostics and action design."],
        ["Tenure and mobility", "YearsAtCompany, YearsInCurrentRole, YearsSinceLastPromotion, YearsWithCurrManager", "Lifecycle trends and career progression."],
        ["Experience", "JobSatisfaction, WorkLifeBalance, EnvironmentSatisfaction, JobInvolvement", "Employee experience distributions and clustering."],
        ["Compensation", "MonthlyIncome, MonthlyRate, PercentSalaryHike, StockOptionLevel", "Income regression and opportunity analysis."],
        ["Personal context", "Age, Gender, MaritalStatus, DistanceFromHome", "Fairness-aware descriptive segmentation."],
        ["Targets", "Attrition, MonthlyIncome", "Classification and regression outcomes."],
    ], [1.35, 2.65, 2.45])
    add_heading(doc, "Data hygiene audit", level=2)
    add_table(doc, ["Audit check", "Result", "Decision"], [
        ["Rows and duplicates", "1,470 rows; zero duplicate records", "Retain all unique employee records."],
        ["Constants", "EmployeeCount=1, StandardHours=80, Over18='Y'", "Drop at load time because they carry no variance."],
        ["Identifier", "EmployeeNumber", "Drop because it is an ID with no predictive meaning."],
        ["Tenure normalization", "Five ordered buckets", "0-1, 2-3, 4-5, 6-10, 11+ in chronological order."],
        ["Leakage control", "Targets excluded from feature frames", "Prevent Attrition and MonthlyIncome from entering their own models."],
    ], [1.55, 2.8, 2.1], PALE_TEAL)
    page_break(doc)

    # Page 5: analytical deep dive.
    add_heading(doc, "Analytical Deep-Dive & Findings")
    add_heading(doc, "Tenure-based lifecycle attrition", level=2)
    add_body(doc, "The tenure chart treats YearsAtCompany as the time axis because this dataset contains no calendar event dates. The bucket order is explicitly categorical and chronological: 0-1, 2-3, 4-5, 6-10, and 11+. This prevents the misleading alphabetical sequence in which 11+ appears before 2-3. The highest observed early-tenure rate is 34.9%, supporting structured onboarding and manager touchpoints during the first year.")
    add_heading(doc, "Driver diagnostics", level=2)
    add_table(doc, ["Signal", "Observed result", "Interpretation"], [
        ["OverTime", "30.5% attrition in the highest group", "Workload and recovery time deserve a policy response."],
        ["JobRole", "39.8% Sales Representative turnover", "Role-specific coaching and progression are warranted."],
        ["MaritalStatus", "25.5% Single attrition", "Use as a context signal, not a causal or individual decision rule."],
        ["DistanceFromHome", "20.7% for >16 km", "Commute burden may amplify retention pressure."],
    ], [1.55, 2.15, 2.75])
    add_heading(doc, "Machine learning integration", level=2)
    add_body(doc, "The attrition workflow uses a stratified Random Forest classifier with balanced class weights. Its validation ROC-AUC is 0.778 and the threshold-selected F1 is 0.522 at an optimal threshold of 0.220. The model quality view also exposes precision, recall, default-threshold F1, and the confusion matrix so stakeholders can understand the trade-off between false alarms and missed at-risk employees.")
    add_body(doc, "The MonthlyIncome workflow uses a tuned Random Forest regressor alongside a Linear Regression baseline. The reported validation result is R² 0.945 with RMSE of $1,091. These models are decision-support components; they do not replace HR judgment, employee conversations, or governance.")
    add_callout(doc, "METHOD LIMITATION", "The dataset is a static historical snapshot. Associations should not be interpreted as causal effects, and model outputs should be monitored for drift, fairness, and changes in workforce composition.", PALE_GOLD)
    page_break(doc)

    # Page 6: visual evidence placeholders 1-2.
    add_heading(doc, "Visual Evidence: Dashboard UI Screenshots")
    add_body(doc, "The following placeholders are intentionally sized and bordered for final dashboard screenshots. Insert captures from the Streamlit application at a consistent desktop viewport so the five-level hierarchy remains legible in review.")
    add_screenshot_placeholder(doc, "[INSERT DASHBOARD SCREENSHOT 1: LEVEL 1 EXECUTIVE WORKFORCE KPIs]")
    add_screenshot_placeholder(doc, "[INSERT DASHBOARD SCREENSHOT 2: LEVEL 2 TENURE LIFECYCLE & INCOME PROGRESSION TRENDS]")
    page_break(doc)

    # Page 7: visual evidence placeholders 3-4 and report export.
    add_heading(doc, "Visual Evidence: Drivers, Risks & Actions")
    add_screenshot_placeholder(doc, "[INSERT DASHBOARD SCREENSHOT 3: LEVEL 3 ATTRITION DRIVERS (OVERTIME & JOB ROLES)]")
    add_screenshot_placeholder(doc, "[INSERT DASHBOARD SCREENSHOT 4: LEVEL 4 & 5 RISKS, ML PERFORMANCE & STRATEGIC ACTIONS]")
    add_callout(doc, "IN-APP REPORT EXPORT", "The Generate Report page produces a current-state PDF containing calculated KPIs, model metrics, and methodological notes. This Word report is the manually authored capstone narrative; the PDF is the live dashboard-state export.", PALE_BLUE)
    page_break(doc)

    # Page 8: actions, learning and references.
    add_heading(doc, "Strategic HR Recommendations & Prioritized Action Matrix")
    add_table(doc, ["Initiative", "Target driver", "Expected HR impact", "Timeline"], [
        ["Overtime workload caps and recovery-time alerts", "30.5% overtime-group attrition", "Reduce sustained workload pressure and improve manager visibility.", "0-90 days pilot; quarterly review"],
        ["Sales Representative coaching and career paths", "39.8% role turnover", "Improve role confidence, progression clarity, and manager support.", "30-120 days"],
        ["30/90/180-day onboarding check-ins", "34.9% early-tenure churn", "Detect adjustment barriers before first-year attrition becomes irreversible.", "Start immediately; monthly cohort review"],
        ["Human-in-the-loop ML risk review", "F1 0.522 at threshold 0.220", "Prioritise supportive conversations while protecting against automated decisions.", "Design in 30 days; audit quarterly"],
    ], [1.7, 1.45, 2.25, 1.1], PALE_TEAL)
    add_heading(doc, "Technical learning outcomes", level=2)
    add_bullets(doc, [
        "pandas: reproducible cleaning, ordered categorical variables, grouped metrics, and feature-frame construction.",
        "Streamlit: cached data and resource functions, multi-level navigation, Plotly integration, formatted dataframes, and PDF download flow.",
        "Plotly: high-contrast chart configuration, readable legends, data labels, and chronological categorical axes.",
        "scikit-learn: preprocessing pipelines, class-balanced modelling, GridSearchCV, cross-validation, threshold tuning, permutation importance, and KMeans.",
        "python-docx: structured report generation, professional styles, tables, borders, shaded callouts, screenshot placeholders, and page fields.",
    ])
    add_heading(doc, "Formal references", level=2)
    add_body(doc, "1. Pavan Subhasht. IBM HR Analytics Employee Attrition & Performance Dataset. Kaggle. https://www.kaggle.com/datasets/pavansubhasht/ibm-hr-analytics-attrition-dataset")
    add_body(doc, "2. IBM SkillsBuild. Data Analytics with AI Internship learning resources, 2026. IBM SkillsBuild in collaboration with BharatCares / CSRBOX.")
    add_body(doc, "3. Pedregosa, F. et al. Scikit-learn: Machine Learning in Python. Journal of Machine Learning Research, 12, 2825-2830, 2011.")
    add_body(doc, "4. McKinney, W. Python for Data Analysis. O'Reilly Media.")
    add_body(doc, "5. Streamlit, Plotly, KaggleHub, ReportLab, and python-docx official documentation.")
    add_callout(doc, "CONCLUSION", "People Signals demonstrates how a compact, production-minded analytics application can move from workforce facts to defensible HR action. Its strongest contribution is the connection between measurable signals, transparent model trade-offs, and human-led retention planning.", PALE_TEAL)

    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build_report())
